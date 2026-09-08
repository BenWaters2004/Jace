from __future__ import annotations

import asyncio
import base64
import io
import json
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image, ImageOps
from pypdf import PdfReader

from jace.attachments.audio import AudioTranscriptionError, transcribe_audio
from jace.attachments.service import attachment_path, update_attachment_processing
from jace.config import settings
from jace.db.models import Attachment


@dataclass
class PreparedAttachment:
    attachment_id: str
    name: str
    kind: str
    context: str = ""
    images: list[str] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


def _trim(text: str, limit: int | None = None) -> str:
    max_chars = limit or settings.attachment_text_max_chars
    cleaned = text.replace("\x00", "").strip()
    if len(cleaned) <= max_chars:
        return cleaned
    return cleaned[:max_chars].rstrip() + "\n…[attachment text truncated]"


def _image_to_base64(path: Path) -> tuple[str, dict]:
    with Image.open(path) as image:
        image = ImageOps.exif_transpose(image)
        if getattr(image, "is_animated", False):
            image.seek(0)
        if image.mode not in {"RGB", "L"}:
            image = image.convert("RGB")
        max_edge = settings.attachment_model_image_max_edge
        if max(image.size) > max_edge:
            image.thumbnail((max_edge, max_edge), Image.Resampling.LANCZOS)
        buffer = io.BytesIO()
        if image.mode == "L":
            image = image.convert("RGB")
        image.save(buffer, format="JPEG", quality=88, optimize=True)
        return base64.b64encode(buffer.getvalue()).decode("ascii"), {
            "width": image.width,
            "height": image.height,
        }


def _extract_pdf_text(path: Path) -> tuple[str, int]:
    reader = PdfReader(str(path))
    pages: list[str] = []
    for index, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            pages.append(f"--- Page {index + 1} ---\n{text.strip()}")
    return _trim("\n\n".join(pages)), len(reader.pages)


def _render_pdf_pages(path: Path, page_limit: int) -> list[str]:
    try:
        import fitz
    except ImportError:
        return []

    images: list[str] = []
    document = fitz.open(str(path))
    try:
        limit = min(page_limit, document.page_count)
        for index in range(limit):
            page = document.load_page(index)
            matrix = fitz.Matrix(settings.attachment_pdf_render_scale, settings.attachment_pdf_render_scale)
            pixmap = page.get_pixmap(matrix=matrix, alpha=False)
            raw = pixmap.tobytes("jpeg")
            images.append(base64.b64encode(raw).decode("ascii"))
    finally:
        document.close()
    return images


def _read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document
        document = Document(str(path))
        parts: list[str] = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return _trim("\n".join(parts))

    raw = path.read_bytes()
    if b"\x00" in raw[:4096]:
        raise ValueError("The document appears to contain unsupported binary data.")
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return _trim(raw.decode(encoding))
        except UnicodeDecodeError:
            continue
    raise ValueError("The document could not be decoded as text.")


async def prepare_attachment(session, attachment: Attachment) -> PreparedAttachment:
    path = attachment_path(attachment)
    if not path.exists():
        raise ValueError(f"Attachment file is missing: {attachment.original_name}")

    if attachment.media_kind == "image":
        image, metadata = _image_to_base64(path)
        return PreparedAttachment(
            attachment_id=attachment.id,
            name=attachment.original_name,
            kind="image",
            context=f"Attached image: {attachment.original_name}",
            images=[image],
            metadata=metadata,
        )

    if attachment.media_kind == "pdf":
        text = attachment.extracted_text
        metadata: dict = {}
        if not text:
            text, page_count = await asyncio.to_thread(_extract_pdf_text, path)
            metadata["page_count"] = page_count
            await update_attachment_processing(session, attachment, extracted_text=text, metadata=metadata)
        else:
            try:
                metadata = json.loads(attachment.metadata_json or "{}")
            except json.JSONDecodeError:
                metadata = {}
        images = await asyncio.to_thread(
            _render_pdf_pages, path, settings.attachment_pdf_render_pages
        )
        context = f"Attached PDF: {attachment.original_name}"
        if text:
            context += f"\n\nExtracted PDF text:\n{text}"
        if images:
            context += f"\n\nThe first {len(images)} PDF page(s) are also attached as images for visual inspection."
        return PreparedAttachment(
            attachment_id=attachment.id,
            name=attachment.original_name,
            kind="pdf",
            context=context,
            images=images,
            metadata=metadata,
        )

    if attachment.media_kind == "document":
        text = attachment.extracted_text
        if not text:
            text = await asyncio.to_thread(_read_document, path)
            await update_attachment_processing(session, attachment, extracted_text=text)
        return PreparedAttachment(
            attachment_id=attachment.id,
            name=attachment.original_name,
            kind="document",
            context=f"Attached document: {attachment.original_name}\n\n{text}",
        )

    if attachment.media_kind == "audio":
        text = attachment.extracted_text
        metadata: dict = {}
        if not text:
            text, metadata = await transcribe_audio(path)
            text = _trim(text)
            await update_attachment_processing(session, attachment, extracted_text=text, metadata=metadata)
        else:
            try:
                metadata = json.loads(attachment.metadata_json or "{}")
            except json.JSONDecodeError:
                metadata = {}
        return PreparedAttachment(
            attachment_id=attachment.id,
            name=attachment.original_name,
            kind="audio",
            context=f"Attached audio: {attachment.original_name}\n\nLocal transcript:\n{text}",
            metadata=metadata,
        )

    raise ValueError(f"Unsupported media kind: {attachment.media_kind}")


async def prepare_attachments(session, attachments: list[Attachment]) -> tuple[str, list[str], list[dict]]:
    contexts: list[str] = []
    images: list[str] = []
    summaries: list[dict] = []
    for attachment in attachments:
        try:
            prepared = await prepare_attachment(session, attachment)
        except AudioTranscriptionError as exc:
            contexts.append(f"Attached audio: {attachment.original_name}\nAudio transcription unavailable: {exc}")
            summaries.append({"id": attachment.id, "name": attachment.original_name, "kind": attachment.media_kind, "error": str(exc)})
            continue
        contexts.append(prepared.context)
        images.extend(prepared.images)
        summaries.append(
            {
                "id": prepared.attachment_id,
                "name": prepared.name,
                "kind": prepared.kind,
                "metadata": prepared.metadata,
            }
        )
    return "\n\n".join(contexts), images, summaries

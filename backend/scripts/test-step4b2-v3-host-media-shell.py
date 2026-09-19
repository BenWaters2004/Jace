from __future__ import annotations

import asyncio
import json
import os
import sys
import tempfile
from pathlib import Path

from PIL import Image
from docx import Document

from jace.attachments.processors import (
    _image_to_base64,
    _read_document,
)
from jace.computer.host_context import (
    host_context_message,
)
from jace.computer.host_request import (
    host_auto_tool_call,
)
from jace.database import SessionLocal
from jace.tools import ensure_tools_registered
from jace.tools.base import ToolContext
from jace.tools.registry import registry
from jace.tools.routing import route_tool_names
from jace.tools.shell_reliability import (
    ShellCommandInput,
    run_shell_command_tool,
)


DOWNLOADS = r"C:\Users\Ben\Downloads"


def continuation_tests() -> None:
    history = [
        {
            "role":
                "user",
            "content":
                (
                    "Show me what's in "
                    + DOWNLOADS
                    + "."
                ),
        },
        {
            "role":
                "assistant",
            "content":
                "Here are the files in Downloads.",
        },
        {
            "role":
                "user",
            "content":
                "Read 12.01.2024.docx",
        },
    ]

    resolved = host_context_message(
        "Read 12.01.2024.docx",
        history,
    )

    expected = (
        DOWNLOADS
        + r"\12.01.2024.docx"
    )

    if expected not in resolved:
        raise RuntimeError(
            f"DOCX follow-up did not inherit Downloads. Got: {resolved!r}"
        )

    image_history = history[:-1] + [
        {
            "role":
                "user",
            "content":
                "Okay describe 16x16.png",
        },
    ]

    image_resolved = host_context_message(
        "Okay describe 16x16.png",
        image_history,
    )

    expected_image = (
        DOWNLOADS
        + r"\16x16.png"
    )

    if expected_image not in image_resolved:
        raise RuntimeError(
            "Image follow-up did not inherit Downloads."
        )

    print("DOCX filename follow-up path inheritance: PASS")
    print("Image filename follow-up path inheritance: PASS")


def routing_tests() -> None:
    docx_message = (
        "Read 12.01.2024.docx\n\n"
        "HOST FILE CONTEXT\n"
        "Resolved path: "
        + DOWNLOADS
        + r"\12.01.2024.docx"
    )

    docx_tools = route_tool_names(
        docx_message
    )

    if "inspect_host_media" not in docx_tools:
        raise RuntimeError(
            "DOCX host request did not route inspect_host_media."
        )

    if "read_host_file" in docx_tools:
        raise RuntimeError(
            "DOCX host request still routed the text-only reader."
        )

    image_message = (
        "Describe 16x16.png\n\n"
        "HOST FILE CONTEXT\n"
        "Resolved path: "
        + DOWNLOADS
        + r"\16x16.png"
    )

    image_tools = route_tool_names(
        image_message
    )

    if "inspect_host_media" not in image_tools:
        raise RuntimeError(
            "Image host request did not route inspect_host_media."
        )

    auto = host_auto_tool_call(
        image_tools,
        image_message,
    )

    if (
        auto is None
        or auto["function"]["name"]
        != "inspect_host_media"
    ):
        raise RuntimeError(
            "Image host request did not build deterministic inspect_host_media call."
        )

    discovery = route_tool_names(
        r"Read a text file from C:\Users\Ben\Downloads."
    )

    if "list_host_directory" not in discovery:
        raise RuntimeError(
            "Directory-level 'read a text file from' request did not list the directory."
        )

    if "read_host_file" in discovery:
        raise RuntimeError(
            "Directory-level read request still routed read_host_file against a directory."
        )

    print("DOCX → inspect_host_media routing: PASS")
    print("Image → inspect_host_media routing: PASS")
    print("Deterministic host media call construction: PASS")
    print("Directory-level file discovery routing: PASS")


def processor_tests() -> None:
    with tempfile.TemporaryDirectory(
        prefix="jace-4b2v3-media-"
    ) as temp:
        root = Path(temp)

        docx_path = (
            root
            / "sample.docx"
        )

        document = Document()
        document.add_paragraph(
            "Jace DOCX extraction test."
        )
        table = document.add_table(
            rows=1,
            cols=2,
        )
        table.cell(
            0,
            0,
        ).text = "Key"
        table.cell(
            0,
            1,
        ).text = "Value"
        document.save(
            str(
                docx_path
            )
        )

        extracted = _read_document(
            docx_path
        )

        if (
            "Jace DOCX extraction test."
            not in extracted
            or "Key | Value"
            not in extracted
        ):
            raise RuntimeError(
                "Existing Jace DOCX processor did not extract expected content."
            )

        image_path = (
            root
            / "16x16.png"
        )

        Image.new(
            "RGB",
            (
                16,
                16,
            ),
            (
                255,
                255,
                255,
            ),
        ).save(
            image_path
        )

        image_b64, metadata = _image_to_base64(
            image_path
        )

        if not image_b64:
            raise RuntimeError(
                "Existing Jace image processor returned no image payload."
            )

        if (
            metadata.get(
                "width"
            )
            != 16
            or metadata.get(
                "height"
            )
            != 16
        ):
            raise RuntimeError(
                "Existing Jace image processor returned wrong dimensions."
            )

    print("Existing DOCX content extraction pipeline: PASS")
    print("Existing image-to-model pipeline: PASS")


def registration_tests() -> None:
    ensure_tools_registered()

    media = registry.get(
        "inspect_host_media"
    )
    shell = registry.get(
        "run_shell_command"
    )

    if media is None:
        raise RuntimeError(
            "inspect_host_media was not registered. Check multimodal settings."
        )

    if (
        media.risk != "read"
        or media.default_permission != "allow"
    ):
        raise RuntimeError(
            "inspect_host_media should be read/Allow."
        )

    if shell is None:
        raise RuntimeError(
            "run_shell_command is not registered."
        )

    if (
        shell.risk != "execute"
        or shell.default_permission != "ask"
    ):
        raise RuntimeError(
            "run_shell_command must remain execute/Ask."
        )

    if (
        getattr(
            shell.handler,
            "__module__",
            "",
        )
        != "jace.tools.shell_reliability"
    ):
        raise RuntimeError(
            "run_shell_command is not using the v3 reliability runner."
        )

    print("Host media tool registration: PASS")
    print("v3 shell runner replaced original shell handler: PASS")
    print("Shell execution remains execute/Ask: PASS")


async def powershell_test() -> None:
    if os.name != "nt":
        print("PowerShell live fixture: SKIP (non-Windows test host)")
        return

    async with SessionLocal() as session:
        context = ToolContext(
            session=session,
            conversation_id=None,
            user_message=(
                "Use PowerShell to output shell-v3-ok."
            ),
        )

        result = await run_shell_command_tool(
            ShellCommandInput(
                shell="powershell",
                command=(
                    "Write-Output 'shell-v3-ok'"
                ),
                timeout_seconds=20,
            ),
            context,
        )

    data = json.loads(
        result.content
    )

    if data[
        "exit_code"
    ] != 0:
        raise RuntimeError(
            "PowerShell fixture failed: "
            + str(
                data
            )
        )

    if (
        "shell-v3-ok"
        not in data[
            "stdout"
        ]
    ):
        raise RuntimeError(
            "PowerShell stdout did not contain the expected fixture text."
        )

    if not data.get(
        "resolved_executable"
    ):
        raise RuntimeError(
            "PowerShell result did not report the resolved executable."
        )

    print(
        "PowerShell executable discovery: PASS "
        f"({data['resolved_executable']})"
    )
    print("PowerShell UTF-8 stdout execution: PASS")


async def main() -> int:
    print()
    print(
        "Jace Step 4B.2 v3 Host Media + Shell Reliability test"
    )
    print(
        "======================================================"
    )
    print()

    continuation_tests()
    routing_tests()
    processor_tests()
    registration_tests()
    await powershell_test()

    print()
    print(
        "PASS - Step 4B.2 v3 host follow-ups, media inspection and shell "
        "reliability are operational."
    )
    print(
        "The test uses temporary media and a harmless PowerShell Write-Output "
        "command only."
    )
    print()

    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(
            asyncio.run(
                main()
            )
        )
    except Exception as exc:
        print()
        print(
            f"FAIL - {exc}",
            file=sys.stderr,
        )
        print()
        raise
